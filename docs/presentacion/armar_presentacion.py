"""
Presentacion de avances del GPR FMCW, sobre la plantilla de FIUBA.

    uv run --with python-pptx python armar_presentacion.py
    (o con cualquier Python que tenga python-pptx)

Sale avances_gpr.pptx en esta carpeta. Las figuras se toman de donde ya
estan en el repo (redaccion/figuras, GPRv2/simulaciones_meep/salidas), asi
que si se regenera una figura alcanza con volver a correr esto.

plantilla_fiuba.pptx son las cinco diapositivas de formato de la facultad:
    0  portada con el logo
    1  titulo del trabajo
    2  separador de seccion (titulo grande abajo)
    3  contenido (titulo arriba)
    4  cierre (www.ingenieria.uba.ar)
Las de seccion y contenido se clonan; la portada, el titulo y el cierre se
usan tal cual.

Ojo con lo que va a cambiar (octubre de 2026): la triangular pasa a salir
de la placa propia del generador, y con ella cambia el sincronismo. Las
diapositivas que lo tocan lo dicen.
"""

import copy
import math
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.normpath(os.path.join(AQUI, "..", ".."))
FIG = os.path.join(RAIZ, "redaccion", "figuras")
MEEP = os.path.join(RAIZ, "GPRv2", "simulaciones_meep", "salidas")
SALIDA = os.path.join(AQUI, "avances_gpr.pptx")

FECHA = "Octubre de 2026"
AUTORES = "Santiago Mogica y Martín Ledesma"

# Colores de la plantilla
AZUL = RGBColor(0x00, 0x95, 0xD8)      # titulos
CELESTE = RGBColor(0x1F, 0xA9, 0xE6)
NARANJA = RGBColor(0xF1, 0x9C, 0x1A)
VERDE = RGBColor(0xA4, 0xC6, 0x39)
BORDO = RGBColor(0xA5, 0x2A, 0x2A)
VIOLETA = RGBColor(0x7B, 0x5A, 0x9C)
GRIS = RGBColor(0x64, 0x74, 0x8B)      # subtitulos
TEXTO = RGBColor(0x1F, 0x29, 0x37)
FONDO_TARJ = RGBColor(0xF3, 0xF6, 0xF9)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FUENTE = "Arial"

# Zona util de una diapositiva de contenido (pulgadas)
X0, X1 = 0.62, 12.85
Y0 = 1.35


# --- clonado de diapositivas de la plantilla -------------------------------

def clonar(prs, idx):
    """Diapositiva nueva con las formas de la de plantilla `idx`.

    Solo copia formas sin relaciones (las franjas de color y los
    placeholders de titulo y numero); las imagenes de la plantilla no se
    copian, que es lo que se quiere.
    """
    base = prs.slides[idx]
    nueva = prs.slides.add_slide(base.slide_layout)
    arbol = nueva.shapes._spTree
    for sh in list(nueva.shapes):
        arbol.remove(sh._element)
    for sh in base.shapes:
        if sh.shape_type == 13:          # imagen
            continue
        if sh.has_text_frame and not sh.is_placeholder and \
                sh.text_frame.text.strip():
            continue                      # cuerpo de texto de ejemplo
        arbol.append(copy.deepcopy(sh._element))
    return nueva


def poner_titulo(slide, texto, tam=None):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.type != 13:   # no sldNum
            tf = sh.text_frame
            p = tf.paragraphs[0]
            for extra in tf.paragraphs[1:]:
                extra._p.getparent().remove(extra._p)
            runs = p.runs
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
            runs[0].text = texto
            if tam:
                runs[0].font.size = Pt(tam)
            return sh
    raise RuntimeError("la plantilla no tiene titulo")


def seccion(prs, titulo, notas=""):
    s = clonar(prs, 2)
    poner_titulo(s, titulo)
    if notas:
        s.notes_slide.notes_text_frame.text = notas
    return s


def contenido(prs, titulo, subtitulo=None, notas=""):
    s = clonar(prs, 3)
    poner_titulo(s, titulo, 32)
    if subtitulo:
        texto(s, X0, 1.12, X1 - X0, 0.45, subtitulo, 20, GRIS)
    if notas:
        s.notes_slide.notes_text_frame.text = notas
    return s


# --- helpers de dibujo -------------------------------------------------------

def texto(slide, x, y, w, h, contenido_, tam=18, color=TEXTO, negrita=False,
          alinear=PP_ALIGN.LEFT, ancla=MSO_ANCHOR.TOP, italica=False):
    """Caja de texto. `contenido_` es un str o una lista de parrafos; cada
    parrafo es un str o una lista de (texto, {opciones}) para mezclar
    estilos en una linea."""
    caja = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = caja.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancla
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    parrafos = contenido_ if isinstance(contenido_, list) else [contenido_]
    for i, par in enumerate(parrafos):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinear
        trozos = par if isinstance(par, list) else [(par, {})]
        for t, op in trozos:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = op.get("fuente", FUENTE)
            f.size = Pt(op.get("tam", tam))
            f.bold = op.get("negrita", negrita)
            f.italic = op.get("italica", italica)
            f.color.rgb = op.get("color", color)
        p.space_after = Pt(6)
    return caja


def vinetas(slide, x, y, w, h, items, tam=20, color=TEXTO, sep=10):
    """Lista con vinetas de verdad (buChar), no un caracter pegado."""
    caja = texto(slide, x, y, w, h, items, tam, color)
    from pptx.oxml.ns import qn
    for p in caja.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.32))))
        pPr.set("indent", str(int(-Inches(0.32))))
        for tag in ("a:buNone", "a:buChar", "a:buFont"):
            for e in pPr.findall(qn(tag)):
                pPr.remove(e)
        bf = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
        bc = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bf)
        pPr.append(bc)
        p.space_after = Pt(sep)
    return caja


def rect(slide, x, y, w, h, relleno, redondeado=True, linea=None):
    forma = MSO_SHAPE.ROUNDED_RECTANGLE if redondeado else MSO_SHAPE.RECTANGLE
    r = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    if redondeado:
        r.adjustments[0] = 0.08
    r.fill.solid()
    r.fill.fore_color.rgb = relleno
    if linea is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = linea
        r.line.width = Pt(1.25)
    r.shadow.inherit = False
    return r


def circulo(slide, x, y, d, relleno, num, tam=16):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                               Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = relleno
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.add_run()
    r.text = str(num)
    r.font.size, r.font.bold, r.font.name = Pt(tam), True, FUENTE
    r.font.color.rgb = BLANCO
    return c


def flecha(slide, x1, y1, x2, y2, color=GRIS, ancho=2.0):
    l = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                   Inches(y1), Inches(x2), Inches(y2))
    l.line.color.rgb = color
    l.line.width = Pt(ancho)
    from pptx.oxml.ns import qn
    ln = l.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med",
                                            "len": "med"})
    ln.append(tail)
    return l


def linea(slide, x1, y1, x2, y2, color=GRIS, ancho=2.0, punteada=False):
    l = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                   Inches(y1), Inches(x2), Inches(y2))
    l.line.color.rgb = color
    l.line.width = Pt(ancho)
    if punteada:
        from pptx.enum.dml import MSO_LINE
        l.line.dash_style = MSO_LINE.DASH
    return l


def imagen(slide, ruta, x, y, w_max, h_max, centrar=True):
    """Imagen escalada para entrar en la caja, sin deformar."""
    with Image.open(ruta) as im:
        wp, hp = im.size
    esc = min(w_max / wp, h_max / hp)
    w, h = wp * esc, hp * esc
    if centrar:
        x += (w_max - w) / 2
        y += (h_max - h) / 2
    return slide.shapes.add_picture(ruta, Inches(x), Inches(y), Inches(w),
                                    Inches(h))


def fuente_fig(slide, x, y, w, txt):
    texto(slide, x, y, w, 0.3, txt, 11, GRIS, italica=True)


def dato(slide, x, y, w, numero, unidad, rotulo, color=AZUL):
    """Numero grande con su rotulo debajo."""
    texto(slide, x, y, w, 0.75,
          [[(numero, {"tam": 36, "negrita": True, "color": color}),
            (unidad, {"tam": 20, "negrita": True, "color": color})]])
    texto(slide, x, y + 0.72, w, 0.6, rotulo, 15, GRIS)


def aviso(slide, x, y, w, h, txt):
    """Recuadro de 'esto va a cambiar'."""
    r = rect(slide, x, y, w, h, RGBColor(0xFF, 0xF4, 0xE0), linea=NARANJA)
    tf = r.text_frame
    tf.word_wrap = True
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.15))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Va a cambiar: "
    r1.font.bold = True
    r2 = p.add_run()
    r2.text = txt
    for rr in (r1, r2):
        rr.font.size, rr.font.name = Pt(15), FUENTE
        rr.font.color.rgb = RGBColor(0x8A, 0x4B, 0x00)
    return r


def tarjeta(slide, x, y, w, h, color, titulo, cuerpo, tam=15, alto_tit=0.5):
    """Tarjeta como las de la plantilla: fondo claro y barra de color arriba.
    `alto_tit` = lugar para el titulo (0.85 si ocupa dos lineas)."""
    rect(slide, x, y, w, h, FONDO_TARJ)
    rect(slide, x, y, w, 0.09, color, redondeado=True)
    texto(slide, x + 0.2, y + 0.25, w - 0.4, alto_tit, titulo, 18, color,
          negrita=True)
    yb = y + 0.25 + alto_tit + 0.15
    texto(slide, x + 0.2, yb, w - 0.4, y + h - yb - 0.1, cuerpo, tam, TEXTO)


# --- el contenido ------------------------------------------------------------

def armar():
    prs = Presentation(os.path.join(AQUI, "plantilla_fiuba.pptx"))

    # 2 - titulo del trabajo (la diapositiva 1 de la plantilla, editada)
    tit = prs.slides[1]
    for sh in tit.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue                      # franjas de color y vacios
        if sh.is_placeholder and sh.placeholder_format.type == 13:
            continue                      # numero de diapositiva
        if sh.is_placeholder:
            poner_titulo(tit, "Radar de penetración terrestre FMCW: "
                              "estado de avance")
        else:
            ps = sh.text_frame.paragraphs
            nuevos = [AUTORES,
                      "Ingeniería Electrónica — Trabajo Práctico Profesional (TPP)",
                      FECHA]
            for p, t in zip(ps, nuevos):
                runs = p.runs
                for r in runs[1:]:
                    r._r.getparent().remove(r._r)
                runs[0].text = t
    tit.notes_slide.notes_text_frame.text = (
        "Presentamos dónde está el radar hoy: qué está armado, cómo se "
        "procesa la señal y qué se midió. Al final, lo que cambia con la "
        "placa del generador propio.")

    # ---------------------------------------------------------------- sistema
    seccion(prs, "El sistema",
            "Primero qué es un GPR FMCW y cómo está armado el nuestro.")

    s = contenido(prs, "Qué mide un radar FMCW",
                  "La distancia se convierte en una frecuencia",
                  "El VCO barre de 1 a 2 GHz con una triangular. Lo que vuelve "
                  "del blanco llega con retardo, y al mezclarlo con lo que sale "
                  "queda un tono, el batido, cuya frecuencia es proporcional a "
                  "la distancia. Medir distancia es medir esa frecuencia con "
                  "una FFT. La resolución la fija el ancho de banda, no el "
                  "muestreo.")
    vinetas(s, X0, 1.85, 5.6, 3.2, [
        "El VCO barre 1–2 GHz con una rampa triangular.",
        "El eco llega retrasado τ = 2d/c: al mezclarlo con la señal que "
        "sale queda un tono de batido.",
        [("f", {"italica": True}), ("batido", {"tam": 13}),
         (" = α₀·τ  →  proporcional a la distancia.", {})],
        "Las rampas de subida y de bajada dan la misma frecuencia.",
    ], 19)
    imagen(s, os.path.join(FIG, "arquitectura", "principio_fmcw.png"),
           6.55, 1.75, 6.3, 3.6)
    dato(s, X0, 5.25, 3.0, "1040", " MHz", "ancho de banda barrido")
    dato(s, 3.95, 5.25, 3.0, "14,4", " cm", "resolución: c / 2B")
    dato(s, 7.3, 5.25, 3.0, "138,6", " Hz/m", "batido por metro (rampa 50 ms)")

    s = contenido(prs, "Diagrama en bloques",
                  notas="La triangular sintoniza el VCO; el splitter da la "
                  "señal a transmitir y el oscilador local. El eco pasa por el "
                  "LNA y el mezclador, se filtra y lo digitaliza un PCM1808 a "
                  "48 kHz. El ESP32-C3 lo manda por USB a la PC junto con "
                  "lecturas de la triangular. Hoy la triangular sale de un "
                  "generador de laboratorio; la placa propia está en camino.")
    imagen(s, os.path.join(FIG, "arquitectura", "diagrama_bloques.png"),
           X0, 1.3, X1 - X0, 4.75)
    aviso(s, X0 + 1.5, 6.15, 9.2, 0.55,
          "la triangular pasa del generador de laboratorio a la placa propia.")

    # ------------------------------------------------------------------ banco
    seccion(prs, "Banco de pruebas",
            "Lo que está armado y las dos cosas que hubo que caracterizar: "
            "el VCO y los cables.")

    s = contenido(prs, "El banco de laboratorio",
                  notas="Cadena de RF sobre una plaqueta, bocinas de chapa "
                  "construidas por nosotros, PCM1808 y ESP32-C3. Los números "
                  "de la derecha son los parámetros con los que se midió todo "
                  "lo que se muestra después.")
    imagen(s, os.path.join(FIG, "banco", "primer_banco_medicion1.jpeg"),
           X0, 1.3, 6.4, 5.2)
    fuente_fig(s, X0, 6.55, 6.4, "Cadena de RF: VCO, atenuador, splitter, "
               "mezclador y amplificadores.")
    filas = [("100", " ms", "período de la triangular: dos rampas de 50 ms"),
             ("48 → 6", " kS/s", "muestreo del PCM1808 y salida diezmada"),
             ("300", "", "muestras por rampa"),
             ("20", " /s", "rampas útiles por segundo (subidas y bajadas)")]
    for i, (n, u, r) in enumerate(filas):
        dato(s, 7.45, 1.35 + i * 1.3, 5.3, n, u, r)

    s = contenido(prs, "El VCO no es lineal",
                  "Su sensibilidad cambia casi 3 a 1 a lo largo de la banda",
                  "Con una rampa de tensión lineal, la frecuencia no crece "
                  "lineal: el batido de un blanco se desparrama y el pico se "
                  "ensancha. La primera solución fue predistorsionar la rampa "
                  "con un DAC. Hoy la rampa es lineal y la corrección se hace "
                  "en el software, remuestreando la señal; se ve en la parte de "
                  "procesamiento.")
    imagen(s, os.path.join(FIG, "vco", "vco_frecuencia_vs_tension.png"),
           X0, 1.75, 7.6, 4.9)
    dato(s, 8.6, 1.9, 4.2, "167–468", " MHz/V", "sensibilidad dF/dV medida")
    dato(s, 8.6, 3.3, 4.2, "943–1982", " MHz", "banda usada, 0 a 3 V")
    vinetas(s, 8.6, 4.7, 4.25, 2.0, [
        "Antes: predistorsión con un DAC.",
        "Ahora: corrección en software, sin tocar el hardware.",
    ], 17)

    s = contenido(prs, "Los cables suman distancia",
                  "Un FMCW mide la diferencia de retardo entre las entradas "
                  "del mezclador",
                  "El cable de antena se suma al retardo del blanco como un "
                  "offset fijo: no cambia la pendiente ni la resolución, solo "
                  "corre el origen. Lo medimos con el VNA, con fase. Cambiar "
                  "los RG-58 viejos por RG-213 de 1 m bajó el offset y la "
                  "pérdida.")
    imagen(s, os.path.join(FIG, "cables", "cables_impacto.png"),
           X0, 1.75, 7.6, 4.9)
    dato(s, 8.6, 1.9, 4.2, "3,86 → 1,45", " m", "offset aparente: RG-58 "
         "viejos → RG-213 nuevos (VNA)")
    dato(s, 8.6, 3.45, 4.2, "4,7", " dB", "de pérdida recuperados")
    vinetas(s, 8.6, 5.0, 4.25, 1.6, [
        "Es un offset, no un error de escala.",
        "Se corrige con un punto de calibración.",
    ], 17)

    # ---------------------------------------------------------- procesamiento
    seccion(prs, "Procesamiento de la señal",
            "El núcleo del trabajo de software: de las muestras a la "
            "distancia.")

    s = contenido(prs, "La cadena de procesamiento",
                  "Cada rampa de 50 ms se procesa sola y después se promedian",
                  "Recorrido completo. En la placa: filtrado y diezmado. En la "
                  "PC: se reconstruye dónde empieza cada rampa, se corta, se "
                  "corrige el VCO, FFT, promedio de varias rampas, fondo y "
                  "pico.")
    pasos = [
        (CELESTE, "Diezmado ×8", "Filtros de semibanda en el ESP32-C3: "
         "48 → 6 kS/s sin aliasing"),
        (NARANJA, "Sincronismo", "Período y fase de la triangular leída "
         "por GPIO3"),
        (VERDE, "Corte de rampas", "300 muestras; las bajadas se invierten "
         "en el tiempo"),
        (BORDO, "Corrección del VCO", "Remuestreo a frecuencia uniforme "
         "(variable θ)"),
        (VIOLETA, "Hann + FFT", "Relleno ×8 (2400 puntos); se guarda |X|"),
        (CELESTE, "Promedio", "N = 8 rampas por fila, incoherente"),
        (NARANJA, "Fondo y detección", "Resta del fondo y puerta de "
         "energía"),
        (VERDE, "Pico sub-bin", "Parábola en dB → distancia"),
    ]
    ancho, alto, gx = 2.75, 2.05, 0.33
    for i, (col, t, d) in enumerate(pasos):
        fila, colum = divmod(i, 4)
        x = X0 + colum * (ancho + gx)
        y = 1.95 + fila * (alto + 0.5)
        rect(s, x, y, ancho, alto, FONDO_TARJ)
        circulo(s, x + 0.18, y + 0.2, 0.5, col, i + 1)
        texto(s, x + 0.8, y + 0.22, ancho - 0.9, 0.5, t, 17, TEXTO,
              negrita=True, ancla=MSO_ANCHOR.MIDDLE)
        texto(s, x + 0.2, y + 0.85, ancho - 0.4, alto - 0.95, d, 15, GRIS)
        if colum < 3:
            flecha(s, x + ancho + 0.03, y + alto / 2, x + ancho + gx - 0.03,
                   y + alto / 2)
    fuente_fig(s, X0, 6.75, 10, "Pasos 1: firmware (adquisicion.ino). "
               "Pasos 2 a 8: PC (vivo_rapido.py).")

    s = contenido(prs, "Sincronismo por software",
                  "Saber en qué muestra empieza cada rampa",
                  "El generador de laboratorio no da sincronismo. Leemos la "
                  "triangular con el ADC del ESP32, unas 19 lecturas por "
                  "período, y ajustamos un triángulo ideal a 50 períodos a la "
                  "vez: período y fase. El error en el borde de rampa queda en "
                  "unos 10 µs, menos de una décima de muestra. Con la placa "
                  "propia esto se reemplaza por la onda cuadrada de "
                  "sincronismo.")
    # triangular dibujada con lineas: dos periodos
    ox, oy, wt, ht = X0 + 0.2, 1.95, 6.4, 2.3
    # vertices: minimo (abajo) en 0, 0.5 y 1; maximo (arriba) en 0.25 y 0.75
    vert = [(ox + px * wt, oy + py * ht) for px, py in
            [(0, 1), (0.25, 0), (0.5, 1), (0.75, 0), (1.0, 1)]]
    for (xa, ya), (xb, yb) in zip(vert, vert[1:]):
        linea(s, xa, ya, xb, yb, NARANJA, 2.5)
    # ~19 lecturas por periodo, con un poco de ruido
    for k in range(37):
        px = 0.012 + k * 0.0271
        fase = (px % 0.5) / 0.5
        ruido = 0.05 * math.sin(k * 7.3)
        cy = oy + abs(1 - 2 * fase) * ht + ruido * ht
        d = 0.09
        c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(ox + px * wt - d / 2),
                               Inches(cy - d / 2), Inches(d), Inches(d))
        c.fill.solid()
        c.fill.fore_color.rgb = BORDO
        c.line.fill.background()
        c.shadow.inherit = False
    for k in range(5):
        xk = ox + k * 0.25 * wt
        linea(s, xk, oy - 0.1, xk, oy + ht + 0.1, GRIS, 1, punteada=True)
    nombres = ["subida", "bajada", "subida", "bajada"]
    for k in range(4):
        xk = ox + k * 0.25 * wt
        col = RGBColor(0xDB, 0xEE, 0xFB) if k % 2 == 0 else \
            RGBColor(0xE6, 0xF2, 0xD3)
        r = rect(s, xk + 0.06, oy + ht + 0.35, 0.25 * wt - 0.12, 0.55, col)
        tf = r.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        rr = p.add_run()
        rr.text = nombres[k] if k % 2 == 0 else "bajada invertida"
        rr.font.size, rr.font.name = Pt(13), FUENTE
        rr.font.color.rgb = TEXTO
    texto(s, ox, oy + ht + 1.0, wt, 0.4,
          "Puntos: lecturas de la triangular (~19 por período). Línea: "
          "triángulo ajustado.", 12, GRIS, italica=True)
    vinetas(s, 7.45, 1.95, 5.4, 3.2, [
        "El ESP32-C3 lee la triangular por GPIO3, 187,5 veces por segundo.",
        "Se ajustan período y fase sobre los últimos 5 s (~50 períodos), "
        "cada 2 s.",
        "Error en el borde de rampa: ~10 µs (una muestra son 167 µs).",
    ], 18)
    aviso(s, 7.45, 5.45, 5.4, 0.95,
          "la placa del generador trae una onda cuadrada de sincronismo, "
          "flanco por flanco.")

    s = contenido(prs, "Corrección del VCO: remuestreo en θ",
                  "En vez de linealizar el barrido, se linealiza el muestreo",
                  "El batido depende de la frecuencia transmitida, no del "
                  "tiempo. Si definimos una variable θ que avanza al ritmo de "
                  "la frecuencia real, en θ el batido vuelve a ser un tono "
                  "puro, sea cual sea la curva del VCO. Como θ es siempre la "
                  "misma, el remuestreo es una matriz fija: un producto por "
                  "rampa.")
    eqs = [
        ("Fase del batido", "Δψ(t) ≈ 2π · τ · f(t)"),
        ("Nueva variable", "θ(t) = [ f(t) − f(0) ] / α₀"),
        ("En θ es un tono puro", "x(θ) = A · cos( 2π · α₀τ · θ + φ )"),
    ]
    for i, (rot, eq) in enumerate(eqs):
        y = 1.9 + i * 1.45
        rect(s, X0, y, 6.6, 1.2, FONDO_TARJ)
        texto(s, X0 + 0.25, y + 0.12, 6.1, 0.35, rot, 15, GRIS, negrita=True)
        texto(s, X0 + 0.25, y + 0.5, 6.1, 0.6,
              [[(eq, {"fuente": "Cambria Math", "tam": 24, "color": AZUL})]])
    tarjeta(s, 7.6, 1.9, 2.5, 3.2, GRIS, "Antes",
            "Predistorsión: el ESP32 escribía con un DAC una rampa deformada "
            "para que la frecuencia saliera recta.")
    tarjeta(s, 10.3, 1.9, 2.55, 3.2, AZUL, "Ahora",
            "Rampa lineal de cualquier generador. La PC interpola la señal a "
            "frecuencia uniforme con la curva medida.")
    texto(s, 7.6, 5.3, 5.25, 0.8, "Requisito: la triangular tiene que ser "
          "lineal y recorrer de 0 a 3 V completos.", 15, GRIS, italica=True)

    s = contenido(prs, "Promedio de rampas: incoherente",
                  "Promediamos magnitudes: el ruido no baja, pero se vuelve "
                  "parejo",
                  "Hay dos formas de promediar. Coherente: sumar los complejos; "
                  "el ruido se cancela y el piso baja 10·log N. Incoherente: "
                  "sumar magnitudes; el piso queda igual pero su dispersión "
                  "baja como raíz de N, y el pico se distingue mejor. Usamos "
                  "incoherente porque no está garantizado que la fase se "
                  "repita entre rampas. Promediar magnitud o potencia da lo "
                  "mismo para ubicar el pico; para medir niveles, la magnitud "
                  "subestima el ruido 1 dB.")
    imagen(s, os.path.join(FIG, "procesamiento", "promedio_vs_N.png"),
           X0, 1.8, 7.9, 3.2)
    fuente_fig(s, X0, 5.05, 7.9, "Monte Carlo con la misma cadena del "
               "programa (docs/promedio_rampas.py).")
    dato(s, 8.85, 1.85, 4.0, "5,3 → 1,6", " dB",
         "dispersión del piso con N = 8")
    dato(s, 8.85, 3.3, 4.0, "+9", " dB", "lo que daría el coherente con N = 8, "
         "si la fase se repite", color=GRIS)
    vinetas(s, X0, 5.55, 12.2, 1.2, [
        "Elegimos incoherente: la estabilidad de fase entre rampas no está "
        "verificada, y las bajadas invertidas tienen otra fase.",
        "Magnitud o potencia: igual para ubicar el pico (< 0,2 dB en "
        "detección); para medir niveles, la magnitud subestima el piso "
        "~1 dB.",
    ], 16, sep=6)

    # ------------------------------------------------------------ resultados
    seccion(prs, "Resultados",
            "Mediciones del 28 de septiembre con una placa metálica a "
            "distintas distancias, contra la simulación en MEEP.")

    s = contenido(prs, "Radargrama en vivo: placa a 90 cm",
                  notas="Arriba el radargrama: el tiempo hacia arriba, la "
                  "frecuencia en horizontal. La traza roja es el pico de cada "
                  "fila. Abajo, el espectro de la última fila sobre el fondo; "
                  "el pico en 498 Hz. 3,59 m aparentes son 0,90 m de aire más "
                  "el offset de cables y bocinas. A la derecha, el estado de "
                  "la cadena y la triangular plegada.")
    imagen(s, os.path.join(FIG, "procesamiento", "radargrama_placa_90cm.png"),
           X0, 1.25, X1 - X0, 5.3)
    fuente_fig(s, X0, 6.6, 12, "vivo_rapido.py, fondo congelado, 8 rampas "
               "por fila (0,4 s). Pico en 498 Hz = 3,59 m aparentes.")

    s = contenido(prs, "Distancia contra frecuencia",
                  "Cinco distancias de la placa: la pendiente es la "
                  "esperada",
                  "Cada punto es una captura. La pendiente medida es 135,7 "
                  "Hz/m contra 138,6 ideal y 139,1 simulada. La ordenada al "
                  "origen es el offset: 1,46 m de cables, medidos con el VNA, "
                  "más 1,23 m de bocinas y electrónica. Medición y simulación "
                  "difieren menos de 6 Hz, o sea menos de 4,5 cm, en todas "
                  "las distancias.")
    imagen(s, os.path.join(MEEP, "mediciones_2026-09-28", "relevamiento.png"),
           X0, 1.75, 8.3, 5.0)
    dato(s, 9.2, 1.85, 3.7, "135,7", " Hz/m", "pendiente medida "
         "(ideal 138,6; MEEP 139,1)")
    dato(s, 9.2, 3.3, 3.7, "2,69", " m", "offset: 1,46 cables + 1,23 bocinas "
         "y electrónica")
    dato(s, 9.2, 4.75, 3.7, "< 4,5", " cm", "diferencia banco − simulación")

    s = contenido(prs, "El banco contra la simulación",
                  "El pico coincide; queda una joroba a la izquierda sin "
                  "explicar",
                  "Espectro medido y simulado de la placa a 90 cm, cada uno "
                  "normalizado a su pico. El pico principal coincide. A la "
                  "izquierda hay una joroba unos 70 Hz abajo que se mueve con "
                  "la placa: no es acoplamiento entre antenas ni la sala. Lo "
                  "próximo es separar subidas de bajadas para ver si viene del "
                  "VCO.")
    imagen(s, os.path.join(MEEP, "placa0.90m_hueco10cm_ancho70cm",
                           "comparacion.png"), X0, 1.75, X1 - X0, 4.3)
    vinetas(s, X0, 6.15, 12.2, 0.8, [
        "Joroba ~70 Hz bajo el pico: se mueve con la placa y persiste con el "
        "fondo restado. Hipótesis a probar: diferencia entre subidas y "
        "bajadas (histéresis del VCO).",
    ], 15)

    # ---------------------------------------------------------------- futuro
    seccion(prs, "Próximos pasos")

    s = contenido(prs, "Qué viene",
                  notas="Lo inmediato es la placa del generador: deja el "
                  "banco sin generador de laboratorio, alimentable desde un "
                  "powerbank, y trae el sincronismo por hardware. Después, "
                  "mejoras del procesamiento y salir del laboratorio.")
    items = [
        (CELESTE, "Generador propio",
         "Placa de la triangular (0–3 V, 20–200 ms) ya diseñada; reemplaza "
         "al generador de laboratorio y permite alimentar todo desde un "
         "powerbank."),
        (NARANJA, "Sincronismo por hardware",
         "La onda cuadrada de la placa marca cada rampa: reemplaza el ajuste "
         "estadístico sobre la triangular."),
        (VERDE, "Procesamiento",
         "Promedio en potencia (niveles sin sesgo); verificar la fase entre "
         "rampas y, si se repite, promedio coherente (+9 dB)."),
        (BORDO, "Joroba de 70 Hz",
         "Comparar subidas contra bajadas para descartar histéresis del "
         "VCO."),
        (VIOLETA, "Fuera del laboratorio",
         "Blancos enterrados y medición sobre suelo real."),
    ]
    w5, g5 = 2.25, 0.2
    for i, (col, t, d) in enumerate(items):
        x = X0 + i * (w5 + g5)
        tarjeta(s, x, 1.55, w5, 4.3, col, t, d, 15, alto_tit=0.85)

    # cierre al final
    lst = prs.slides._sldIdLst
    cierre = lst[4]
    lst.remove(cierre)
    lst.append(cierre)
    # las diapositivas 2 y 3 de la plantilla eran solo moldes
    for idx in (3, 2):
        sid = lst[idx]
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)

    prs.save(SALIDA)
    print("Generado:", SALIDA, "-", len(prs.slides), "diapositivas")


if __name__ == "__main__":
    armar()
